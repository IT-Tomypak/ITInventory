"""Merge FAR sheet CE + PC Masterlist_2026 + IT Accessories into one sheet using CE's column layout."""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FAR = 'TFP FAR 30 JUNE 2025 (FY2025).xlsx'
PC = 'PC List_Latest_2026 (1).xlsx'
ACC = 'IT Accessories list.xlsx'
OUT = 'IT Asset Merged.xlsx'

# CE columns first, then the inventory details the FAR lacks
HEAD = ['Source', 'Acquisition Date', 'Disposal Date', 'F/A Code', 'F/A Description', 'Suppliers',
        'Cost Ctr. 1', 'Cost Ctr. 2', 'Cost Ctr. 3', 'Purchase Order', 'NO.', 'Invoice No.', 'Quantity',
        'Cost c/f', 'NBV As At 30/06/2025', 'Remarks',
        'Asset Tag', 'Category', 'Brand', 'Model', 'Serial No', 'User', 'Employee ID', 'Status', 'OS', 'RAM', 'Storage']
COL = {h: i for i, h in enumerate(HEAD)}


def row(**kw):
    r = [None] * len(HEAD)
    for k, v in kw.items():
        r[COL[k]] = v
    return r


def load(path, sheet):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    keys = next(it)
    return [dict(zip(keys, r)) for r in it]


# --- PC Masterlist, keyed by tag (TPPC218) so FAR rows can join onto it
pcs = {}
for p in load(PC, 'Masterlist_2026'):
    tag = p['Device_name']
    if not tag or tag == 'abc':
        continue
    pcs[tag] = p


def pc_fields(p):
    return dict(**{'Asset Tag': p['Device_name'], 'Category': p['Device_type'], 'Brand': p['Brand'],
                   'Model': p['Model'], 'Serial No': p['Serial_number'], 'User': p['Employee_name'],
                   'Employee ID': p['Employee_id'], 'Status': p['Status'], 'OS': p['OS'],
                   'RAM': p['RAM'], 'Storage': p['HDD_capacity']})


out = []

# --- FAR CE: data starts at row 6, stops at first blank F/A Code
ws = openpyxl.load_workbook(FAR, data_only=True)['CE']
for r in ws.iter_rows(min_row=6, values_only=True):
    if not r[2]:
        break
    rec = row(**{'Source': 'FAR (CE)', 'Acquisition Date': r[0], 'Disposal Date': r[1], 'F/A Code': r[2],
                 'F/A Description': r[3], 'Suppliers': r[4], 'Cost Ctr. 1': r[5], 'Cost Ctr. 2': r[6],
                 'Cost Ctr. 3': r[7], 'Purchase Order': r[8], 'NO.': r[9], 'Invoice No.': r[10],
                 'Quantity': r[11], 'Remarks': r[13], 'Cost c/f': r[15], 'NBV As At 30/06/2025': r[38]})
    m = re.search(r'TPPC-?(\d+)', str(r[3]))
    p = pcs.pop(f'TPPC{m.group(1)}', None) if m else None
    if p:
        rec[COL['Source']] = 'FAR (CE) + PC List'
        for k, v in pc_fields(p).items():
            rec[COL[k]] = v
    out.append(rec)

# --- remaining PCs not in the FAR
for p in pcs.values():
    out.append(row(**{'Source': 'PC List', 'Acquisition Date': p['Purchased_year'],
                      'F/A Description': ' '.join(filter(None, [p['Brand'], p['Model']])),
                      'Suppliers': p['Vendor_name'], 'Cost Ctr. 1': p['Plant'], 'Cost Ctr. 2': p['Department'],
                      'Quantity': 1, 'Remarks': p['Remarks']}, **pc_fields(p)))

# --- IT accessories
for a in load(ACC, 'Master'):
    if not a['Details']:
        continue
    out.append(row(**{'Source': 'IT Accessories', 'F/A Description': a['Details'], 'Suppliers': a['Supplier'],
                      'Cost Ctr. 1': a['OwnerShip2'], 'Cost Ctr. 2': a['Department'], 'Quantity': 1,
                      'Remarks': a['Remark'], 'Category': a['Categories'], 'Model': a['Model'],
                      'Serial No': a['Serial_no'], 'User': a['OwnerShip'], 'Employee ID': a['Employee_no'],
                      'Status': a['Status']}))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'IT (CE merged)'
ws.append(HEAD)
for r in out:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='305496')
for col in (COL['Acquisition Date'], COL['Disposal Date']):
    for c in ws[get_column_letter(col + 1)][1:]:
        c.number_format = 'dd/mm/yyyy'
for col in (COL['Cost c/f'], COL['NBV As At 30/06/2025']):
    for c in ws[get_column_letter(col + 1)][1:]:
        c.number_format = '#,##0.00'
for i, h in enumerate(HEAD, 1):
    ws.column_dimensions[get_column_letter(i)].width = 45 if h == 'F/A Description' else max(12, len(h) + 2)
ws.freeze_panes = 'B2'
ws.auto_filter.ref = ws.dimensions
wb.save(OUT)

from collections import Counter
print(len(out), Counter(r[0] for r in out))
